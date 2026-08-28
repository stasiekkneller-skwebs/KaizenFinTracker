from django.shortcuts import render
from django.views.generic.edit import FormView
from django.views.generic.base import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views import View
from django.views.generic import ListView, DetailView
from django.views.generic.edit import UpdateView, DeleteView
from .forms import CategoryCreationForm, TransactionForm, CategoryUpdateForm, BudgetForm, BudgetUpdateForm
from django.contrib import messages
from .models import Category, Transaction, Budget
from django.shortcuts import get_object_or_404
from django.db.models import Sum
import csv
import io
from django.http import JsonResponse
from datetime import datetime
import os
from decimal import Decimal
from django.conf import settings
from django.views import View
from django.http import HttpResponse, Http404
from django.template.loader import render_to_string
from django.contrib.auth.mixins import LoginRequiredMixin
from django.utils import timezone
from django.utils import translation
import openpyxl
from datetime import date

# Create your views here.
class HomeView(TemplateView):
    template_name = 'tracker/home.html'

class UploadFileView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    template_name = 'tracker/upload.html'

    def get(self, request):
        return render(request, self.template_name)

    def post(self, request):
        file = request.FILES.get('file')

        if not file:
            return JsonResponse({'status': 'error', 'message': 'Brak pliku'}, status=400)

        imported = 0
        skipped = 0

        try:
            decoded = file.read().decode('utf-8')
            reader = csv.DictReader(io.StringIO(decoded))
            imported_transactions = []
            for row in reader:
                try:
                    category, _ = Category.objects.get_or_create(
                        name=row['kategoria'],
                        user=request.user
                    )
                    Transaction.objects.create(
                        user=request.user,
                        category=category,
                        name = row['nazwa'],
                        amount=row['kwota'],
                        transaction_type=row['typ'],
                        description=row['opis'],
                        transaction_date=row['data']
                    )
                    imported_transactions.append({
    'nazwa': row['nazwa'],
    'kwota': row['kwota'],
    'kategoria': row['kategoria'],
    'typ': row['typ'],
    'data': row['data']
})
                    imported += 1
                except Exception:
                    skipped += 1

        except Exception as e:
            print(f"Błąd wiersza: {e}, dane: {row}")
            skipped += 1

        return JsonResponse({'status': 'ok', 'imported': imported, 'skipped': skipped, 'transactions': imported_transactions})

class DashboardView(LoginRequiredMixin, TemplateView):
    login_url = '/user/login/'
    template_name = 'tracker/dashboard.html'

    def _get_curr_year_and_month(self):
        year = datetime.now().year
        month = datetime.now().month 
        return year, month
    
    def get_queryset(self):
        year, month = self._get_curr_year_and_month()
        return Transaction.objects.filter(user=self.request.user, transaction_date__month=month,
            transaction_date__year=year)

    

    def _get_sum_of_transactions(self):
        qs = self.get_queryset()
        expense = qs.filter(transaction_type='expense').aggregate(total=Sum('amount'))['total'] or 0
        income = qs.filter(transaction_type='income').aggregate(total=Sum('amount'))['total'] or 0
        return expense, income
    
    def _get_budgets(self):
        year, month = self._get_curr_year_and_month()
        budgets = Budget.objects.filter(user=self.request.user, month=month,
            year=year)
        return budgets
    
    def _get_over_categories(self, expenses_data):
        cat_names = []
        budgets = self._get_budgets()
        for ed in expenses_data:
            for b in budgets:
                if ed['category__name'] == b.category.name and ed['total'] > b.limit:
                    cat_names.append(b.category.name)
        return cat_names
            
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['categories'] = Category.objects.filter(user=self.request.user)
        expenses, incomes = self._get_sum_of_transactions()
        context['total_expenses'] = -expenses
        context['total_incomes'] = incomes
        context['balance'] = incomes - expenses
        context['budgets'] = self._get_budgets()
        expenses_data = self.get_queryset().filter(
    transaction_type='expense'
).values('category__name').annotate(total=Sum('amount'))
        context['expenses_data'] = expenses_data
        context['over_categories'] = self._get_over_categories(expenses_data)
        context['chart_labels'] = list(expenses_data.values_list('category__name', flat=True))
        context['chart_data'] = list(expenses_data.values_list('total', flat=True))
        context['chart_colors'] = list(expenses_data.values_list('category__color', flat=True))
        return context
    
    

class AddCategoryView(LoginRequiredMixin, FormView):
    login_url = '/user/login/'
    form_class = CategoryCreationForm
    success_url = '/categories/'
    template_name = 'tracker/add_category.html'

    def form_valid(self, form):
        category = form.save(commit=False)
        category.user = self.request.user
        category.save()
        messages.success(self.request, 'Dodano nową kategorię')
        return super().form_valid(form)
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

class AddTransactionView(LoginRequiredMixin, FormView):
    login_url = '/user/login/'
    form_class = TransactionForm
    success_url = '/transactions/'
    template_name = 'tracker/add_transaction.html'

    def _get_curr_year_and_month(self):
        year = datetime.now().year
        month = datetime.now().month 
        return year, month

    def form_valid(self, form):
        transaction = form.save(commit=False)
        transaction.user = self.request.user
        transaction.save()

        # sprawdź czy budżet przekroczony
        self._check_budget_exceeded(transaction)

        messages.success(self.request, 'Dodano transakcję')
        return super().form_valid(form)

    def _check_budget_exceeded(self, transaction):
        if transaction.transaction_type != 'expense':
            return
        
        year, month = transaction.transaction_date.year, transaction.transaction_date.month
        budget = Budget.objects.filter(
            user=self.request.user,
            category=transaction.category,
            month=month,
            year=year
        ).first()

        if not budget:
            return

        spent = Transaction.objects.filter(
            user=self.request.user,
            category=transaction.category,
            transaction_type='expense',
            transaction_date__month=month,
            transaction_date__year=year
        ).aggregate(total=Sum('amount'))['total'] or 0

        if spent > budget.limit:
            spent=float(spent)
            limit=float(budget.limit)
            messages.error(self.request, f'Przekroczono limit {limit}, wydano {spent}  ')
            
    
    
    
    


class CategoryListView(LoginRequiredMixin, ListView):
    login_url = '/user/login/'
    model = Category
    template_name = 'tracker/category_list.html'
    context_object_name = 'categories'

    def get_queryset(self):
        return Category.objects.filter(user=self.request.user)
    
class TransactionListView(LoginRequiredMixin, ListView):
    login_url = '/user/login/'
    model = Transaction
    template_name = 'tracker/transaction_list.html'
    context_object_name = 'transactions'
    paginate_by = 20

    def get_queryset(self):
        qs = Transaction.objects.filter(user=self.request.user)
        
        transaction_type = self.request.GET.get('type')
        category = self.request.GET.get('category')
        year = self.request.GET.get('year')
        month = self.request.GET.get('month')

        if transaction_type:
            qs = qs.filter(transaction_type=transaction_type)
        if category:
            qs = qs.filter(category__id=category)
        if year:
            qs = qs.filter(transaction_date__year=year)
            if month:
                qs = qs.filter(transaction_date__month=month)

        return qs
    
    def _get_sum_of_transactions(self):
        qs = self.get_queryset()
        expense = qs.filter(transaction_type='expense').aggregate(total=Sum('amount'))['total'] or 0
        income = qs.filter(transaction_type='income').aggregate(total=Sum('amount'))['total'] or 0
        return expense, income
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['categories'] = Category.objects.filter(user=self.request.user)
        expenses, incomes = self._get_sum_of_transactions()
        context['total_expenses'] = -expenses
        context['total_incomes'] = incomes
        context['balance'] = incomes - expenses
        return context

class ExportTransactionsView(LoginRequiredMixin, View):
    login_url = '/user/login/'



    def _get_transactions(self, request):
        qs = Transaction.objects.filter(user=self.request.user)
                
        transaction_type = self.request.GET.get('type')
        category = self.request.GET.get('category')
        year = self.request.GET.get('year')
        month = self.request.GET.get('month')
        
        if transaction_type:
            qs = qs.filter(transaction_type=transaction_type)
        if category:
            qs = qs.filter(category__id=category)
        if year:
            qs = qs.filter(transaction_date__year=year)
            if month:
                qs = qs.filter(transaction_date__month=month)
        return qs

    def get(self, request):
        transactions = self._get_transactions(request)
        wb = openpyxl.Workbook()
        sheet = wb.active
        headers = ['Data', 'Nazwa', 'Kategoria', 'Kwota', 'Typ', 'Opis']
        for col_num, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=col_num, value=header)
        for row_num, transaction in enumerate(transactions, start=2):
            sheet.cell(row=row_num, column=1, value=str(transaction.transaction_date))
            sheet.cell(row=row_num, column=2, value=transaction.name)
            sheet.cell(row=row_num, column=3, value=str(transaction.category) if transaction.category else 'Brak')
            sheet.cell(row=row_num, column=4, value=float(transaction.amount))
            sheet.cell(row=row_num, column=5, value=transaction.get_transaction_type_display())
            sheet.cell(row=row_num, column=6, value=transaction.description)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        
        response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
        response['Content-Disposition'] = f'attachment; filename="transakcje_{date.today()}.xlsx"'
        return response

    """ def get(self, request):
        format = request.GET.get('format', 'xlsx')
        if format == 'csv':
            return self._export_csv(request)
        return self._export_xlsx(request) """

class CategoryDetailView(LoginRequiredMixin, DetailView):
    login_url = '/user/login/'
    model = Category
    template_name = 'tracker/category_details.html'
    context_object_name = 'category'

    def get_object(self):
        return get_object_or_404(Category, pk=self.kwargs['pk'], user=self.request.user)

class TransactionDetailView(LoginRequiredMixin, DetailView):
    login_url = '/user/login/'
    model = Transaction
    template_name = 'tracker/transaction_details.html'
    context_object_name = 'transaction'

    def get_object(self):
        return get_object_or_404(Transaction, pk=self.kwargs['pk'], user=self.request.user)


class CategoryUpdateView(LoginRequiredMixin, UpdateView):
    model = Category
    form_class = CategoryUpdateForm
    template_name = 'tracker/edit_category.html'
    success_url = '/categories/'
    login_url = '/user/login/'

    def get_object(self):
        return get_object_or_404(Category, pk=self.kwargs['pk'], user=self.request.user)

    def form_valid(self, form):
        messages.success(self.request, 'Zaktualizowano kategorię')
        return super().form_valid(form)
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs
    
class TransactionUpdateView(LoginRequiredMixin, UpdateView):
    model = Transaction
    form_class = TransactionForm
    template_name = 'tracker/edit_transaction.html'
    success_url = '/transactions/'
    login_url = '/user/login/'

    def get_object(self):
        return get_object_or_404(Transaction, pk=self.kwargs['pk'], user=self.request.user)

    def form_valid(self, form):
        messages.success(self.request, 'Zaktualizowano transakcję')
        return super().form_valid(form)
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs
    

class CategoryDeleteView(LoginRequiredMixin, DeleteView):
    login_url = '/user/login/'
    model = Category
    success_url = '/categories/'
    template_name = 'tracker/category_delete.html'
    context_object_name = 'category'

    def get_object(self):
        return get_object_or_404(Category, pk=self.kwargs['pk'], user=self.request.user)

class TransactionDeleteView(LoginRequiredMixin, DeleteView):
    login_url = '/user/login/'
    model = Transaction
    success_url = '/transactions/'
    template_name = 'tracker/transaction_delete.html'
    context_object_name = 'transaction'

    def get_object(self):
        return get_object_or_404(Transaction, pk=self.kwargs['pk'], user=self.request.user)

class BudgetDeleteView(LoginRequiredMixin, DeleteView):
    login_url = '/user/login/'
    model = Budget
    success_url = '/budgets/'
    template_name = 'tracker/budget_delete.html'
    context_object_name = 'budgets'

    def get_object(self):
        return get_object_or_404(Budget, pk=self.kwargs['pk'], user=self.request.user)

class AddBudgetView(LoginRequiredMixin, FormView):
    login_url = '/user/login/'
    model = Budget
    form_class = BudgetForm
    success_url = '/budgets/'
    template_name = 'tracker/add_budget.html'

    def form_valid(self, form):
        transaction = form.save(commit=False)
        transaction.user = self.request.user
        transaction.save()
        messages.success(self.request, 'Ustalono nowy budżet')
        return super().form_valid(form)
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

class UpdateBudgetView(LoginRequiredMixin, UpdateView):
    model = Budget
    form_class = BudgetUpdateForm
    template_name = 'tracker/edit_budget.html'
    success_url = '/budgets/'
    login_url = '/user/login/'

    def get_object(self):
        return get_object_or_404(Budget, pk=self.kwargs['pk'], user=self.request.user)

    def form_valid(self, form):
        messages.success(self.request, 'Zaktualizowano budżet')
        return super().form_valid(form)
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

class BudgetListView(LoginRequiredMixin, ListView):
    login_url = '/user/login/'
    model = Budget
    template_name = 'tracker/budget_list.html'
    context_object_name = 'budgets'
    paginate_by = 20

    def _get_filters_data(self):
        budget_year = self.request.GET.get('year')
        budget_month = self.request.GET.get('month') 
        return budget_year, budget_month

    def get_queryset(self):
        qs = Budget.objects.filter(user=self.request.user)
        budget_year, budget_month = self._get_filters_data()
        if budget_year and budget_month:
            qs = qs.filter(month = budget_month).filter(year = budget_year)
        elif budget_year:
            qs = qs.filter(year = budget_year)
        return qs
    
    def _get_month_name(self, num_of_month):
        month_names = {
    1: "styczeń",
    2: "luty",
    3: "marzec",
    4: "kwiecień",
    5: "maj",
    6: "czerwiec",
    7: "lipiec",
    8: "sierpień",
    9: "wrzesień",
    10: "październik",
    11: "listopad",
    12: "grudzień",
}
        try:
            key = int(num_of_month)
        except (ValueError, TypeError):
            return ''  # Jeśli nie da się zmienić na int (np. self._get_month_name('abc'))

        # Metoda .get(klucz, wartosc_domyslna) załatwia sprawę
        return month_names.get(key, '')
    

    def _get_spent_for_budget(self, budget):
        spent = Transaction.objects.filter(
            user=self.request.user,
            category=budget.category,
            transaction_type='expense',
            transaction_date__month=budget.month,
            transaction_date__year=budget.year
        ).aggregate(total=Sum('amount'))['total'] or 0
        return spent
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        budgets_with_progress = []

        for budget in context['budgets']:
            spent = self._get_spent_for_budget(budget)
            percent = (spent / budget.limit * 100) if budget.limit else 0
            budgets_with_progress.append({
                'budget': budget,
                'spent': spent,
                'percent': round(percent, 1),
                'exceeded': percent > 100
            })

        context['budgets_with_progress'] = budgets_with_progress
        context['year']= self._get_filters_data()[0]
        context['month'] = self._get_month_name(self._get_filters_data()[1])
        return context


